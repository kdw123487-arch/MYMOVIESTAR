"""Poster colour analysis for My Movie Star.

For every film with a TMDB poster (data/films.json, field tm.p):
  1. download the poster and reduce it to ~10 colour clusters,
  2. pick 제1색 (the colour that defines the poster: large area, favouring colour over plain black/white)
     and 제2색 (the strongest accent that is clearly different from 제1색),
  3. turn 제1색 into the film-page background: same hue, darkened so body text stays readable,
  4. turn 제2색 into the point colour: same hue, lifted until it contrasts with the background
     (WCAG 7:1, and 4.5:1 for the dimmer label tone the page mixes from it).

Writes 포스터1색 / 포스터2색 (raw) and 배경색 / 포인트색 (usable) into data/films.xlsx,
and a preview sheet data/poster_colors.jpg. Run tools/xlsx_to_json.py afterwards.
Needs: pip install openpyxl pillow
"""
import io, json, math, sys, time, urllib.request
from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
XLSX, JSON, PREVIEW = ROOT / "data" / "films.xlsx", ROOT / "data" / "films.json", ROOT / "data" / "poster_colors.jpg"
IMG = "https://image.tmdb.org/t/p/w185"
IVORY, MUTED = (241, 232, 214), (168, 154, 140)


# ---------- colour maths ----------
def lin(c):
    c /= 255
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def unlin(x):
    x = 12.92 * x if x <= 0.0031308 else 1.055 * x ** (1 / 2.4) - 0.055
    return x * 255


def lum(rgb):
    r, g, b = (lin(v) for v in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = lum(a), lum(b)
    return (max(la, lb) + 0.05) / (min(la, lb) + 0.05)


def to_oklch(rgb):
    r, g, b = (lin(v) for v in rgb)
    l = (0.4122214708 * r + 0.5363325363 * g + 0.0514459929 * b) ** (1 / 3)
    m = (0.2119034982 * r + 0.6806995451 * g + 0.1073969566 * b) ** (1 / 3)
    s = (0.0883024619 * r + 0.2817188376 * g + 0.6299787005 * b) ** (1 / 3)
    L = 0.2104542553 * l + 0.7936177850 * m - 0.0040720468 * s
    A = 1.9779984951 * l - 2.4285922050 * m + 0.4505937099 * s
    B = 0.0259040371 * l + 0.7827717662 * m - 0.8086757660 * s
    return L, math.hypot(A, B), math.degrees(math.atan2(B, A)) % 360


def from_oklch(L, C, H):
    """OKLCH -> sRGB, lowering chroma until the colour fits in sRGB."""
    while True:
        A, B = C * math.cos(math.radians(H)), C * math.sin(math.radians(H))
        l = (L + 0.3963377774 * A + 0.2158037573 * B) ** 3
        m = (L - 0.1055613458 * A - 0.0638541728 * B) ** 3
        s = (L - 0.0894841775 * A - 1.2914855480 * B) ** 3
        rgb = (4.0767416621 * l - 3.3077115913 * m + 0.2309699292 * s,
               -1.2684380046 * l + 2.6097574011 * m - 0.3413193965 * s,
               -0.0041960863 * l - 0.7034186147 * m + 1.7076147010 * s)
        if all(-1e-4 <= v <= 1 + 1e-4 for v in rgb) or C <= 0.001:
            return tuple(max(0, min(255, round(unlin(max(0.0, min(1.0, v)))))) for v in rgb)
        C -= 0.004


def oklab_dist(a, b):
    La, Ca, Ha = to_oklch(a)
    Lb, Cb, Hb = to_oklch(b)
    ax, ay = Ca * math.cos(math.radians(Ha)), Ca * math.sin(math.radians(Ha))
    bx, by = Cb * math.cos(math.radians(Hb)), Cb * math.sin(math.radians(Hb))
    return math.sqrt((La - Lb) ** 2 + (ax - bx) ** 2 + (ay - by) ** 2)


def hexc(rgb):
    return "#%02x%02x%02x" % tuple(rgb)


def mix(a, b, t):
    return tuple(round(x * t + y * (1 - t)) for x, y in zip(a, b))


# ---------- poster analysis ----------
def clusters(data):
    im = Image.open(io.BytesIO(data)).convert("RGB")
    w, h = im.size
    im = im.crop((int(w * .03), int(h * .03), int(w * .97), int(h * .97)))   # drop scan edges
    im.thumbnail((90, 135))
    q = im.quantize(colors=10, method=Image.Quantize.MEDIANCUT, kmeans=3)
    pal, counts = q.getpalette(), q.getcolors()
    total = sum(c for c, _ in counts)
    return [(c / total, tuple(pal[i * 3:i * 3 + 3])) for c, i in sorted(counts, reverse=True)], im


def pick(cl):
    def vivid(rgb):
        L, C, _ = to_oklch(rgb)
        return C, L
    def score1(share, rgb):
        C, L = vivid(rgb)
        extreme = 0.35 if L < 0.18 or L > 0.95 else 1.0      # plain black / white rarely "is" the poster's colour
        return share * (0.35 + 6 * C) * extreme
    first = max(cl, key=lambda x: score1(*x))[1]
    def score2(share, rgb):
        d = oklab_dist(rgb, first)
        if d < 0.07:
            return -1
        C, L = vivid(rgb)
        C1, _, H1 = to_oklch(first)
        _, _, H = to_oklch(rgb)
        hue = 1.0
        if C >= 0.04 and C1 >= 0.04:                       # both colourful: prefer a different hue, not a shade
            dh = abs((H - H1 + 180) % 360 - 180)
            hue = 0.45 + 0.55 * min(1.0, dh / 60)
        return math.sqrt(share) * (0.12 + 6 * C) * min(1.0, d / 0.18) * hue
    rest = [x for x in cl if x[1] != first]
    second = max(rest, key=lambda x: score2(*x))[1] if rest else first
    return first, second


def background(first):
    L, C, H = to_oklch(first)
    C = min(C * 0.85, 0.075)
    for Lb in [x / 1000 for x in range(300, 150, -5)]:
        bg = from_oklch(Lb, C, H)
        if lum(bg) <= 0.03 and contrast(IVORY, bg) >= 11 and contrast(MUTED, bg) >= 4.5:
            return bg
    return from_oklch(0.16, C, H)


def point(second, bg):
    L, C, H = to_oklch(second)
    if C >= 0.035:
        C = min(max(C, 0.10), 0.17)
    for Lp in [x / 1000 for x in range(max(680, round(L * 1000)), 951, 5)]:
        ac = from_oklch(Lp, C, H)
        if contrast(ac, bg) >= 7 and contrast(mix(ac, bg, 0.7), bg) >= 4.5:
            return ac
    return from_oklch(0.92, min(C, 0.06), H)


def fetch(path):
    for attempt in range(3):
        try:
            with urllib.request.urlopen(IMG + path, timeout=20) as r:
                return r.read()
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2)


# ---------- spreadsheet + preview ----------
def column(ws, name):
    head = [c.value for c in ws[1]]
    if name in head:
        return head.index(name) + 1
    col = ws.max_column + 1
    cell = ws.cell(1, col, name)
    ref = ws.cell(1, col - 1)
    cell.font, cell.fill = copy(ref.font), copy(ref.fill)
    ws.column_dimensions[cell.column_letter].width = 11
    return col


def main():
    films = {f["en"]: f for f in json.loads(JSON.read_text(encoding="utf-8"))["films"]}
    wb = load_workbook(XLSX)
    ws = wb["추천작"]
    head = [c.value for c in ws[1]]
    c_en = head.index("원제") + 1
    c_r1, c_r2 = column(ws, "포스터1색"), column(ws, "포스터2색")
    c_bg, c_ac = column(ws, "배경색"), column(ws, "포인트색")

    rows, done, skipped = [], 0, []
    for r in range(2, ws.max_row + 1):
        en = ws.cell(r, c_en).value
        f = films.get(en)
        p = ((f or {}).get("tm") or {}).get("p")
        if not p:
            skipped.append(en)
            continue
        try:
            cl, thumb = clusters(fetch(p))
        except Exception as e:
            skipped.append(f"{en} ({e})")
            continue
        first, second = pick(cl)
        bg = background(first)
        ac = point(second, bg)
        for col, rgb in ((c_r1, first), (c_r2, second)):
            cell = ws.cell(r, col, hexc(rgb))
            cell.fill = PatternFill("solid", fgColor=hexc(rgb)[1:])
            cell.font = Font(color="000000" if lum(rgb) > 0.35 else "FFFFFF")
        cb = ws.cell(r, c_bg, hexc(bg))
        cb.fill, cb.font = PatternFill("solid", fgColor=hexc(bg)[1:]), Font(color="F1E8D6")
        ca = ws.cell(r, c_ac, hexc(ac))
        ca.fill, ca.font = PatternFill("solid", fgColor=hexc(bg)[1:]), Font(color=hexc(ac)[1:], bold=True)
        rows.append((en, thumb, cl, first, second, bg, ac))
        done += 1
        print(f"{en:40s} 1:{hexc(first)} 2:{hexc(second)} -> bg {hexc(bg)}  point {hexc(ac)}  ({contrast(ac, bg):.1f}:1)")
    wb.save(XLSX)

    # preview: poster | its colour clusters | 제1색 제2색 | the film page swatch
    W, RH, COLS = 600, 150, 2
    sheet = Image.new("RGB", (W * COLS, RH * math.ceil(len(rows) / COLS)), (20, 10, 13))
    d = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 15)
        small = ImageFont.truetype("DejaVuSans.ttf", 12)
    except OSError:
        font = small = ImageFont.load_default()
    for i, (en, thumb, cl, first, second, bg, ac) in enumerate(rows):
        x, y = (i % COLS) * W, (i // COLS) * RH
        t = thumb.copy()
        t.thumbnail((90, 135))
        sheet.paste(t, (x + 8, y + 8))
        bx = x + 108
        for share, rgb in cl:                                   # cluster strip, width = share
            w = max(2, round(share * 200))
            d.rectangle((bx, y + 10, bx + w, y + 30), fill=rgb)
            bx += w
        d.rectangle((x + 108, y + 40, x + 158, y + 70), fill=first)
        d.rectangle((x + 162, y + 40, x + 212, y + 70), fill=second)
        d.text((x + 108, y + 74), "1st   2nd", font=small, fill=(150, 140, 130))
        d.rectangle((x + 330, y + 10, x + 590, y + 140), fill=bg)
        d.text((x + 342, y + 22), en[:26], font=font, fill=ac)
        d.text((x + 342, y + 50), "Synopsis", font=font, fill=ac)
        d.text((x + 342, y + 76), "Body text on the page", font=small, fill=IVORY)
        d.text((x + 342, y + 112), f"{hexc(bg)}  {hexc(ac)}", font=small, fill=mix(ac, bg, 0.7))
    sheet.save(PREVIEW, quality=86, optimize=True)
    print(f"\n{done}편 분석 · 건너뜀 {len(skipped)}편{': ' + ', '.join(skipped) if skipped else ''}")
    if done == 0:
        sys.exit("포스터를 하나도 분석하지 못했습니다.")


if __name__ == "__main__":
    main()
