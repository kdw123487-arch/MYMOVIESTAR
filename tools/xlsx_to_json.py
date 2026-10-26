"""films.xlsx -> films.json converter for My Movie Star.

Usage:  python tools/xlsx_to_json.py            (reads data/films.xlsx, writes data/films.json)
Needs:  pip install openpyxl

Sheet "추천작": one film per row, columns named in COLUMNS below (order doesn't matter).
Sheet "제외작": films removed from the list; kept so earlier ballots still show their titles.
"""
import json, re, sys, unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = ROOT / "data" / "films.xlsx", ROOT / "data" / "films.json"

# Excel header -> JSON field
COLUMNS = {
    "키": "key", "제목": "ko", "원제": "en", "연도": "y", "국가": "n", "장르": "g",
    "러닝타임": "rt", "감독": "dir", "출연": "cast", "수상": "aw", "한줄소개": "p",
    "줄거리": "s", "추천이유": "w", "관람포인트1": "pt1", "관람포인트2": "pt2", "관람포인트3": "pt3",
    "넷플릭스오리지널": "nf", "TMDB_ID": "tmdb", "관람등급": "age", "색상": "pal", "문양": "m", "영문폰트": "ef", "배경색": "bg",
}
MOTIFS = {"sun", "waves", "bands", "rings", "split", "peaks", "arch", "rays", "grid", "dots", "trees"}
FONTS = {"classic", "elegant", "noir", "epic", "typewriter", "mono", "bold", "scifi", "gothic", "western", "retro", "script", "hand", "comic", "light", "neon"}
AGES = {"전체": "ALL", "ALL": "ALL", "12": "12", "15": "15", "18": "18", "19": "18", "청불": "18", "청소년관람불가": "18"}


def slug(en):
    s = unicodedata.normalize("NFD", en.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return "r:" + re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if any(v not in (None, "") for v in r):
            yield {head[i]: v for i, v in enumerate(r) if i < len(head) and head[i]}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def film(row, n):
    f = {COLUMNS[k]: clean(v) for k, v in row.items() if k in COLUMNS}
    for need in ("ko", "en"):
        if not f.get(need):
            sys.exit(f"{n}행: '{[k for k, v in COLUMNS.items() if v == need][0]}' 칸이 비어 있습니다.")
    out = {"key": f.get("key") or slug(f["en"]), "ko": f["ko"], "en": f["en"]}
    for k in ("y", "n", "g", "dir", "cast", "aw", "p", "s", "w"):
        if f.get(k):
            out[k] = f[k]
    for k in ("rt", "tmdb"):
        if f.get(k):
            out[k] = int(float(f[k]))
    pts = [f[k] for k in ("pt1", "pt2", "pt3") if f.get(k)]
    if pts:
        out["pt"] = pts
    if f.get("nf", "").upper() in ("O", "Y", "1", "TRUE", "예"):
        out["nf"] = 1
    if f.get("age"):
        a = AGES.get(f["age"].replace("세", "").strip())
        if not a:
            sys.exit(f"{n}행: 관람등급은 전체/12/15/18 중 하나여야 합니다 ({f['age']}).")
        out["age"] = a
    if f.get("pal"):
        pal = [c.strip() for c in f["pal"].split(",")]
        if len(pal) != 3 or not all(re.fullmatch(r"#[0-9a-fA-F]{6}", c) for c in pal):
            sys.exit(f"{n}행: 색상은 '#aabbcc, #aabbcc, #aabbcc' 형식이어야 합니다.")
        out["pal"] = pal
    if f.get("m"):
        if f["m"] not in MOTIFS:
            sys.exit(f"{n}행: 문양은 {', '.join(sorted(MOTIFS))} 중 하나여야 합니다.")
        out["m"] = f["m"]
    if f.get("ef"):
        if f["ef"] not in FONTS:
            sys.exit(f"{n}행: 영문폰트는 {', '.join(sorted(FONTS))} 중 하나여야 합니다.")
        out["ef"] = f["ef"]
    if f.get("bg"):
        if not re.fullmatch(r"#[0-9a-fA-F]{6}", f["bg"]):
            sys.exit(f"{n}행: 배경색은 #aabbcc 형식이어야 합니다.")
        out["bg"] = f["bg"]
    return out


def main():
    wb = load_workbook(SRC, read_only=True, data_only=True)
    films = [film(r, i + 2) for i, r in enumerate(rows(wb["추천작"]))]
    keys = [f["key"] for f in films]
    dup = {k for k in keys if keys.count(k) > 1}
    if dup:
        sys.exit("키가 겹칩니다: " + ", ".join(sorted(dup)))
    retired = {}
    if "제외작" in wb.sheetnames:
        for r in rows(wb["제외작"]):
            f = {COLUMNS[k]: clean(v) for k, v in r.items() if k in COLUMNS}
            k = f.get("key") or slug(f["en"])
            retired[k] = {x: f[x] for x in ("ko", "en", "y", "n") if f.get(x)}
    OUT.write_text(json.dumps({"films": films, "retired": retired}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(films)}편 → {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
